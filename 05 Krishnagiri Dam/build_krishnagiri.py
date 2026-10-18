"""Adopted EAC for Krishnagiri Dam (KRP): Copernicus DEM above its 480.5 m water surface + power-law storage below.

python build_krishnagiri.py COP_DEM.tif SRTM_DEM.tif
DEMs: Copernicus GLO-30 / SRTM 1" clips of tile N12E078 covering 78.06-78.30 E, 12.40-12.62 N (see README.md).
"""
import os, sys, csv, numpy as np
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'scripts'))
from eac_from_dem import load_dem, eac_table
OUT = HERE + os.sep
POLY = os.path.join(HERE, 'krishnagiri_area_490m.geojson')
COP, SRTM = sys.argv[1], sys.argv[2]
FT, MCFT = 0.3048, 0.0283168466  # m per ft, MCM per Mcft
FRL, MWL, FLAT, TOP = 483.11, 484.63, 480.5, 490.0
GATE = 6.10                      # radial gate height (8 Nos 12.19 x 6.10 m)
CREST = FRL - GATE               # ogee crest, derived
GZ = FRL - 52 * FT               # gauge zero: FRL reads 52 ft on the WRD gauge
S_PUB = 1666 * MCFT              # WRD full capacity at 52 ft
S_1958 = 68.2                    # original gross storage, MCM
# WRD daily bulletin readings (gauge ft, storage Mcft)
BULLETIN = [('01 Sep 2026', 46.85, 1129), ('04 Oct 2026', 50.20, 1463), ('07 Oct 2026', 50.35, 1479),
            ('18 Oct 2026', 50.00, 1442), ('Full (52 ft)', 52.00, 1666)]
zc, ca, crs, _ = load_dem(COP, POLY)
zs, ca_s, _, _ = load_dem(SRTM, POLY)
def dem(z, c, L):  # area km2, storage MCM from a DEM at level L
    r = eac_table(z, c, L, L + 1, 1)[0]; return r[5], r[6]
A_W = dem(zc, ca, FLAT + 1e-6)[0]
def hidden(total):  # storage at the DEM water surface and power-law exponent for S(FRL) = total
    s_w = total - dem(zc, ca, FRL)[1]; return s_w, A_W * (FLAT - GZ) / s_w
S_W, B = hidden(S_PUB)
S_W58, B58 = hidden(S_1958)
def adopted(h, s_w=S_W, b=B):
    if h <= GZ: return 0.0, 0.0, 'below gauge zero'
    if h <= FLAT:
        x = (h - GZ) / (FLAT - GZ)
        return s_w * b / (FLAT - GZ) * x ** (b - 1), s_w * x ** b, 'power law (below DEM water surface)'
    a, s = dem(zc, ca, h); return a, s_w + s, 'Copernicus DEM + storage below 480.5 m'
# datum check: storage increments between bulletin readings above the DEM water surface vs the DEM
def misfit(d):
    pts = [(GZ + f * FT + d, m * MCFT) for _, f, m in BULLETIN]
    return np.mean([abs((dem(zc, ca, h1)[1] - dem(zc, ca, h0)[1]) - (s1 - s0)) for (h0, s0), (h1, s1) in zip(pts, pts[1:]) if h1 > h0])
offsets = np.round(np.arange(-2, 2.01, 0.1), 2)
best = offsets[np.argmin([misfit(d) for d in offsets])]
levels = sorted(set([round(GZ, 3)] + [round(x, 2) for x in np.arange(468.0, TOP + 0.01, 0.5)]
                    + [FLAT, round(CREST, 2), FRL, MWL]))
rows = []
for h in levels:
    a, s, src = adopted(h)
    rows.append([h, round((h - GZ) / FT, 2), int(round(a * 1e6 / ca)), a * 1e6, s * 1e6, a, s, s / MCFT,
                 adopted(h, S_W58, B58)[1], src])
key = [('Gauge zero (FRL - 52 ft)', GZ), ('Spillway crest (FRL - gate 6.10 m)', CREST), ('DEM water surface (Copernicus)', FLAT),
       ('FRL', FRL), ('MWL', MWL), ('Top of curve', TOP)]
from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook(); ws = wb.active; ws.title = 'EAC adopted'
ws.append(['Elevation (m)', 'Gauge (ft)', 'Pixels (30 m)', 'Area (m2)', 'Capacity (m3)', 'Area (km2)', 'Capacity (MCM)',
           'Capacity (Mcft)', 'Capacity, 1958 variant (MCM)', 'Source'])
for r in rows: ws.append(r)
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = 'A2'
for col, w in zip('ABCDEFGHIJ', (13, 11, 12, 14, 16, 11, 14, 15, 26, 40)): ws.column_dimensions[col].width = w
for col in 'DE':
    for c in ws[col][1:]: c.number_format = '#,##0'
for col in 'FGHI':
    for c in ws[col][1:]: c.number_format = '0.000' if col == 'F' else '0.00'
k = wb.create_sheet('Key levels')
k.append(['Level', 'Elevation (m)', 'Gauge (ft)', 'Area (km2)', 'Capacity (MCM)', 'Capacity (Mcft)', '1958 variant (MCM)', 'SRTM DEM from 483 m (MCM)'])
for name, h in key:
    a, s, _ = adopted(h)
    k.append([name, round(h, 2), round((h - GZ) / FT, 2), round(a, 2), round(s, 2), round(s / MCFT, 0),
              round(adopted(h, S_W58, B58)[1], 2), round(dem(zs, ca_s, h)[1], 2) if h > 483 else None])
for c in k[1]: c.font = Font(bold=True)
for col, w in zip('ABCDEFGH', (36, 14, 11, 11, 15, 15, 18, 24)): k.column_dimensions[col].width = w
chk = wb.create_sheet('WRD bulletin check')
chk.append(['Date', 'Gauge (ft)', 'Elevation (m)', 'WRD storage (Mcft)', 'Adopted curve (Mcft)', 'Difference (%)'])
for d, f, m in BULLETIN:
    s = adopted(GZ + f * FT)[1] / MCFT; chk.append([d, f, round(GZ + f * FT, 2), m, round(s, 0), round(100 * (s - m) / m, 1)])
for c in chk[1]: c.font = Font(bold=True)
for col in 'ABCDEF': chk.column_dimensions[col].width = 20
raw = wb.create_sheet('Copernicus raw')
raw.append(['BaseLevel_m', 'Pixels', 'Area_m2', 'Capacity_m3', 'Area_ha', 'Area_km2', 'Capacity_MCM', 'Area_acre', 'Capacity_acre_ft'])
for r in eac_table(zc, ca, 478.0, TOP, 0.5): raw.append(r)
rs = wb.create_sheet('SRTM raw')
rs.append([c.value for c in raw[1]])
for r in eac_table(zs, ca_s, 479.0, TOP, 0.5): rs.append(r)
notes = wb.create_sheet('Notes')
for line in [
 ['Reservoir', 'Krishnagiri Dam (KRP), Thenpennaiyar, Krishnagiri district (dam 12.4753 N, 78.1853 E)'],
 ['Prepared', '18 Oct 2026 with the area-elevation-capacity skill; planning level'],
 ['DEM', f'Copernicus GLO-30 (2011-2015, EGM2008), reprojected to {crs} at 30 m; SRTM 1" (Feb 2000) as cross-check'],
 ['Polygon', 'krishnagiri_area_490m.geojson: cells below 490 m connected to the lake, dam axis (dam_axis.txt) burnt as a wall, traced on the DEM crest from the left-bank hill to the right-bank hill'],
 ['Basin closure', 'With the dam axis the basin holds to 501.0 m, then spills over a natural saddle near 12.522 N, 78.199 E. Without the abutment extension the fill leaks round the left (south-west) end at 492.3 m'],
 ['Below 480.5 m', f'Copernicus flattens the lake at 480.5 m ({A_W:.2f} km2), so storage below it is invisible. Filled with a power law from gauge zero {GZ:.2f} m: S = {S_W:.2f} MCM x ((h-{GZ:.2f})/{FLAT-GZ:.2f})^{B:.2f}, with S(480.5) chosen so S(FRL) = WRD full capacity 1,666 Mcft ({S_PUB:.2f} MCM)'],
 ['1958 variant', f'Same, forced to the original gross storage 68.2 MCM: S(480.5) = {S_W58:.2f} MCM, exponent {B58:.2f}. Use it to bound the dam-break storage'],
 ['Levels', f'FRL {FRL} m (52 ft), MWL {MWL} m (published). Gauge zero {GZ:.2f} m and crest {CREST:.2f} m are derived (FRL - 52 ft; FRL - 6.10 m gate height), confirm from drawings. TBL not found'],
 ['Datum', f'Copernicus heights are EGM2008, WRD levels GTS/MSL. WRD bulletin storage increments above 480.5 m fit the DEM best at an offset of {best:+.1f} m (DEM = WRD + offset); no offset applied'],
 ['Capacity', '1,666 Mcft is the capacity WRD now uses in its daily bulletin; the 1958 gross storage was 68.2 MCM (2,408 Mcft). The difference is probably sedimentation; a bathymetric survey would replace the power-law part'],
 ['SRTM', 'SRTM (Feb 2000) flattens the lake at 483 m and reads about 1.7 m higher than Copernicus on land; shown for comparison only'],
 ['Sources', 'WRD/TNAU reservoir bulletin (52 ft, 1,666 Mcft), Oneindia daily dam levels (Sep-Oct 2026), Wikipedia / IJIRSET 2019 case study (FRL 483.11 m, MWL 484.63 m, 68.2 MCM, 8 gates 12.19 x 6.10 m)'],
]: notes.append(line)
notes.column_dimensions['A'].width = 16; notes.column_dimensions['B'].width = 150
wb.save(OUT + 'Krishnagiri_EAC.xlsx')
# fine table for level-pool routing (hydrology-dambreak engine.py reads this format)
with open(OUT + 'Krishnagiri_EAC_0.05m.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['elevation_m', 'area_m2', 'storage_m3'])
    for h in np.round(np.arange(GZ, TOP + 1e-9, 0.05), 3):
        a, s, _ = adopted(h); w.writerow([h, round(a * 1e6, 1), round(s * 1e6, 1)])
# chart
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
H = np.arange(GZ, TOP + 0.001, 0.05); AD = np.array([adopted(h)[:2] for h in H])
fig, ax = plt.subplots(figsize=(10, 7))
ax.plot(AD[:, 1], H, color='#1f5aa6', lw=2.5, label='Capacity – adopted (MCM)')
ax.plot([adopted(h, S_W58, B58)[1] for h in H], H, '--', color='#1f5aa6', alpha=.6, lw=1.2, label='Capacity – 1958 variant (68.2 MCM at FRL)')
Hs = H[H > 483]; ax.plot([dem(zs, ca_s, h)[1] for h in Hs], Hs, ':', color='grey', lw=1.2, label='Capacity – SRTM from 483 m (not used)')
ax.plot([m * MCFT for _, _, m in BULLETIN], [GZ + f * FT for _, f, _ in BULLETIN], 'o', color='k', ms=4, label='WRD bulletin readings')
ax.set_xlabel('Capacity (MCM)', color='#1f5aa6'); ax.set_ylabel('Water level (m)'); ax.set_xlim(0, 160); ax.set_ylim(466, 491); ax.grid(alpha=.3)
top = ax.twiny()
top.plot(AD[:, 0], H, color='#c8551b', lw=2.5, label='Area – adopted (km²)')
top.set_xlim(32, 0); top.set_xlabel('Area (km²)', color='#c8551b')
for name, h, c in (('Gauge zero', GZ, 'k'), ('Crest', CREST, 'k'), ('DEM water surface', FLAT, 'grey'), ('FRL', FRL, 'g'), ('MWL', MWL, 'orange')):
    ax.axhline(h, color=c, lw=0.8, ls='-.'); ax.text(158, h + 0.1, f'{name} {h:.2f} m', ha='right', fontsize=8, color=c)
sec = ax.secondary_yaxis('right', functions=(lambda h: (h - GZ) / FT, lambda f: GZ + f * FT)); sec.set_ylabel('WRD gauge (ft)')
h1, l1 = ax.get_legend_handles_labels(); h2, l2 = top.get_legend_handles_labels()
fig.legend(h1 + h2, l1 + l2, loc='lower center', ncol=3, fontsize=8, frameon=False)
fig.suptitle('Krishnagiri Dam – Elevation-Area-Capacity Curves (Copernicus GLO-30, planning level)')
fig.tight_layout(rect=(0, 0.08, 1, 1)); fig.savefig(OUT + 'Krishnagiri_EAC.png', dpi=150)
print(f'A_W {A_W:.3f} km2  S_W {S_W:.2f} MCM  b {B:.3f}  | 1958: S_W {S_W58:.2f} b {B58:.3f}  | best datum offset {best:+.1f} m')
for name, h in key:
    a, s, _ = adopted(h)
    print(f'{name:36s} {h:7.2f} m {(h-GZ)/FT:6.2f} ft {a:6.2f} km2 {s:7.2f} MCM {s/MCFT:6.0f} Mcft  1958 {adopted(h, S_W58, B58)[1]:6.2f}')
for d, f, m in BULLETIN:
    s = adopted(GZ + f * FT)[1] / MCFT; print(f'{d:14s} {f:5.2f} ft  WRD {m:5d}  curve {s:5.0f} Mcft  {100*(s-m)/m:+.1f}%')
