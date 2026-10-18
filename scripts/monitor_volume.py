"""Reservoir volume and water level from a surface-water-area time series and an EAC table.

Same calculation as the template's Calcs sheet: each observed area (km2) is linearly interpolated
in the EAC table to a volume (MCM) and a water level (m). Areas outside the table's range are
reported as "Area below/above Calibration range" instead of being extrapolated.

Usage:
  python monitor_volume.py EAC.xlsx|csv AREA.csv --out monitoring.xlsx [--plot monitoring.png]
                           [--name "Mirani Reservoir"] [--start 2006-01-01] [--drop 2025-06-22 ...]

EAC: output of eac_from_dem.py, the QGIS script's export (BaseLevel_m, Area_km2/Area_m2, Capacity_MCM/
Capacity_m3), or a filled template workbook (its Data sheet is read).
AREA: Global Water Watch "Surface Water Area (#ID).csv" (time, Surface Water Area (km2)), or any
CSV/XLSX with a date column and an area-in-km2 column.
"""
import argparse
import csv
import sys
from datetime import date, datetime

import numpy as np

M2_PER_ACRE = 4046.8564224
M3_PER_ACRE_FT = 1233.4818375475
FT_PER_M = 3.280839895


def _rows(path, sheet=None):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [r for r in csv.reader(f)]


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _find_header(rows, wanted):
    """Index of the first row containing all `wanted` predicates, and the matched column indices."""
    for i, r in enumerate(rows[:30]):
        names = [str(c).strip().lower() if c is not None else "" for c in r]
        cols = []
        for pred in wanted:
            j = next((k for k, n in enumerate(names) if pred(n)), None)
            if j is None:
                break
            cols.append(j)
        else:
            return i, cols
    return None, None


def load_eac(path):
    """Return (elev m, area km2, volume MCM) arrays sorted by elevation."""
    is_template = False
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        is_template = "Data" in load_workbook(path, read_only=True).sheetnames
    rows = _rows(path, "Data" if is_template else None)
    if is_template:
        # Template Data sheet: elevation, pixels, Area (m2), Capacity (m3) in columns C-F.
        i, cols = _find_header(rows, [lambda n: "elevation" in n, lambda n: n == "area (m2)",
                                      lambda n: n == "capacity (m3)"])
        scale = (1e-6, 1e-6)
    else:
        i, cols = _find_header(rows, [lambda n: n.startswith("baselevel") or "elev" in n,
                                      lambda n: n == "area_km2", lambda n: n == "capacity_mcm"])
        scale = (1.0, 1.0)
        if i is None:
            i, cols = _find_header(rows, [lambda n: n.startswith("baselevel") or "elev" in n,
                                          lambda n: n == "area_m2", lambda n: n == "capacity_m3"])
            scale = (1e-6, 1e-6)
    if i is None:
        sys.exit(f"Could not find elevation/area/capacity columns in {path}")
    data = []
    for r in rows[i + 1:]:
        vals = [_num(r[c]) if c < len(r) else None for c in cols]
        if None not in vals:
            data.append((vals[0], vals[1] * scale[0], vals[2] * scale[1]))
    if len(data) < 2:
        sys.exit(f"Fewer than two EAC rows in {path}")
    a = np.array(sorted(data))
    return a[:, 0], a[:, 1], a[:, 2]


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def load_series(path):
    rows = _rows(path)
    i, cols = _find_header(rows, [lambda n: n in ("time", "date") or "date" in n,
                                  lambda n: "area" in n])
    if i is None:
        sys.exit(f"Could not find time/date and area columns in {path}")
    out = []
    for r in rows[i + 1:]:
        d = _parse_date(r[cols[0]]) if cols[0] < len(r) else None
        v = _num(r[cols[1]]) if cols[1] < len(r) else None
        if d and v is not None:
            out.append((d, v))
    out.sort()
    return out


def interp_like_template(x, xs, ys):
    """Template's MATCH(...,1)/INDEX interpolation: last row with xs <= x, then linear to the next."""
    if x < xs.min():
        return "Area below Calibration range"
    if x > xs.max():
        return "Area above Calibration range"
    k = int(np.searchsorted(xs, x, side="right")) - 1
    if k >= len(xs) - 1 or xs[k] == xs[k + 1]:
        return float(ys[k])
    return float(ys[k] + (x - xs[k]) * (ys[k + 1] - ys[k]) / (xs[k + 1] - xs[k]))


def suspect_flags(areas, window=2, tol=0.25):
    """Flag points that jump more than tol from the median of their neighbours (cloud/classifier spikes)."""
    a = np.asarray(areas, float)
    flags = []
    for i in range(len(a)):
        nb = np.concatenate([a[max(0, i - window):i], a[i + 1:i + 1 + window]])
        med = np.median(nb) if nb.size else a[i]
        flags.append(bool(med > 0 and abs(a[i] - med) / med > tol))
    return flags


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("eac")
    p.add_argument("area")
    p.add_argument("--out", required=True, help=".xlsx or .csv")
    p.add_argument("--plot", help="PNG with volume and level over time")
    p.add_argument("--name", default="Reservoir")
    p.add_argument("--start", help="drop observations before this date (e.g. before first filling)")
    p.add_argument("--drop", nargs="*", default=[], help="dates to remove (cloudy/bad images)")
    a = p.parse_args()

    elev, area, vol = load_eac(a.eac)
    order = np.argsort(area, kind="stable")  # MATCH needs areas ascending
    area_s, vol_s, elev_s = area[order], vol[order], elev[order]
    series = load_series(a.area)
    start = _parse_date(a.start) if a.start else None
    drop = {_parse_date(d) for d in a.drop}
    series = [(d, v) for d, v in series if (not start or d >= start) and d not in drop]
    if not series:
        sys.exit("No observations left after filtering")
    flags = suspect_flags([v for _, v in series])

    header = ["No", "Date", "Area (km2)", "Volume (MCM)", "Elevation (m)", "Area (Acre)",
              "Volume (Acre-ft)", "Elevation (ft)", "Suspect spike"]
    out = []
    for n, ((d, ak), flag) in enumerate(zip(series, flags), 1):
        v = interp_like_template(ak, area_s, vol_s)
        z = interp_like_template(ak, area_s, elev_s)
        ok = isinstance(v, float)
        out.append([n, d, ak, v, z, ak * 1e6 / M2_PER_ACRE,
                    v * 1e6 / M3_PER_ACRE_FT if ok else None,
                    z * FT_PER_M if ok else None, "check" if flag else ""])

    annual = {}
    for r in out:
        if isinstance(r[3], float):
            annual.setdefault(r[1].year, []).append(r[3])
    annual_rows = [[y, min(v), sum(v) / len(v), max(v), len(v)] for y, v in sorted(annual.items())]

    if a.out.lower().endswith(".csv"):
        with open(a.out, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(out)
    else:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Monitoring"
        ws.append(header)
        for r in out:
            ws.append(r)
            ws.cell(ws.max_row, 2).number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16
        ys = wb.create_sheet("Annual")
        ys.append(["Year", "Min Volume (MCM)", "Average Volume (MCM)", "Max Volume (MCM)", "Observations"])
        for r in annual_rows:
            ys.append(r)
        wb.save(a.out)

    if a.plot:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        good = [r for r in out if isinstance(r[3], float)]
        sus = [r for r in good if r[8]]
        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(11, 7), sharex=True)
        ax1.plot([r[1] for r in good], [r[3] for r in good], color="#1f5aa6", lw=1.2, marker=".", ms=3)
        ax1.plot([r[1] for r in sus], [r[3] for r in sus], "o", mfc="none", color="#c8551b",
                 label="suspect spike – check image")
        ax1.set_ylabel("Volume (MCM)")
        ax1.grid(True, alpha=0.3)
        if sus:
            ax1.legend(loc="upper left")
        ax2.plot([r[1] for r in good], [r[4] for r in good], color="#2c8a5a", lw=1.2, marker=".", ms=3)
        ax2.set_ylabel("Water level (m, DEM datum)")
        ax2.grid(True, alpha=0.3)
        fig.suptitle(f"{a.name} – storage and water level from surface water area")
        fig.tight_layout()
        fig.savefig(a.plot, dpi=150)

    outside = sum(1 for r in out if not isinstance(r[3], float))
    print(f"{len(out)} observations, {outside} outside the EAC range, "
          f"{sum(1 for r in out if r[8])} flagged as suspect -> {a.out}")
    if out and isinstance(out[-1][3], float):
        print(f"Latest {out[-1][1]}: area {out[-1][2]:.2f} km2, volume {out[-1][3]:.2f} MCM, "
              f"level {out[-1][4]:.2f} m")


if __name__ == "__main__":
    main()
