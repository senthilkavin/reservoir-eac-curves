"""Adopted EAC for Chembarambakkam Tank: Copernicus DEM above its 21.0 m water surface + power-law storage below.

python build_chembarambakkam.py COP_DEM.tif SRTM_DEM.tif [hydrology-dambreak elevation_storage_discharge.csv]
DEMs: Copernicus GLO-30 / SRTM 1" mosaics covering 79.85-80.15 E, 12.90-13.12 N (see README.md).
"""
import os, sys, csv, numpy as np
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'scripts'))
from eac_from_dem import load_dem, eac_table
OUT = HERE + os.sep
POLY = os.path.join(HERE, 'chembarambakkam_area_28.5m.geojson')
COP, SRTM = sys.argv[1], sys.argv[2]
STUDY = sys.argv[3] if len(sys.argv) > 3 else os.path.expanduser(
    '~/hydrology-dambreak/examples/chembarambakkam/HEC-HMS_Chembarambakkam/elevation_storage_discharge.csv')
GZ, FLAT, FRL, MWL, TBL, SILL, S_PUB = 18.71472, 21.0, 26.02992, 26.94432, 28.77312, 23.02992, 3645 * 0.0283168466
MCFT = 0.0283168466  # MCM per Mcft
zc, ca, crs, _ = load_dem(COP, POLY)
zs, ca_s, _, _ = load_dem(SRTM, POLY)
def dem(z, c, L):  # area km2, storage MCM from a DEM at level L
    r = eac_table(z, c, L, L + 1, 1)[0]; return r[1], r[5], r[6]
A21 = dem(zc, ca, FLAT + 1e-6)[1]
S21 = S_PUB - dem(zc, ca, FRL)[2]
b = A21 * (FLAT - GZ) / S21
def adopted(h):
    if h <= GZ: return 0.0, 0.0, 'below gauge zero'
    if h <= FLAT:
        x = (h - GZ) / (FLAT - GZ)
        return S21 * b / (FLAT - GZ) * x ** (b - 1), S21 * x ** b, 'power law (below DEM water surface)'
    _, a, s = dem(zc, ca, h); return a, S21 + s, 'Copernicus DEM + storage below 21.0 m'
# study's assumed curve (hydrology-dambreak)
st = list(csv.DictReader(open(STUDY)))
st_h = np.array([float(r['elevation_m']) for r in st]); st_s = np.array([float(r['storage_1000m3']) for r in st]) / 1000
st_a = np.array([float(r['area_km2']) for r in st])
levels = sorted(set([round(GZ, 3)] + [round(x, 2) for x in np.arange(19.0, 28.51, 0.25)] + [21.0, round(SILL, 3), round(FRL, 3), round(MWL, 3)]))
rows = []
for h in levels:
    a, s, src = adopted(h)
    pix = int(round(a * 1e6 / ca))
    rows.append([h, round(h * 3.280839895, 2), pix, a * 1e6, s * 1e6, a, s, s / MCFT, src])
key = [('Gauge zero', GZ), ('DEM water surface (Copernicus)', FLAT), ('Regulator sill', SILL), ('FRL', FRL), ('MWL', MWL), ('Top of DEM-closed basin', 28.5)]
from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook(); ws = wb.active; ws.title = 'EAC adopted'
ws.append(['Elevation (m)', 'Elevation (ft)', 'Pixels (30 m)', 'Area (m2)', 'Capacity (m3)', 'Area (km2)', 'Capacity (MCM)', 'Capacity (Mcft)', 'Source'])
for r in rows: ws.append(r)
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = 'A2'
for col, w in zip('ABCDEFGHI', (13, 13, 12, 14, 16, 11, 14, 15, 40)): ws.column_dimensions[col].width = w
for col in 'DE':
    for c in ws[col][1:]: c.number_format = '#,##0'
for col in 'FGH':
    for c in ws[col][1:]: c.number_format = '0.000' if col == 'F' else '0.00'
k = wb.create_sheet('Key levels')
k.append(['Level', 'Elevation (m)', 'Area (km2)', 'Capacity (MCM)', 'Capacity (Mcft)', 'Study assumed curve (MCM)', 'SRTM DEM from 19 m (MCM)'])
for name, h in key:
    a, s, _ = adopted(h)
    k.append([name, round(h, 3), round(a, 2), round(s, 2), round(s / MCFT, 0), round(float(np.interp(h, st_h, st_s)), 2),
              round(dem(zs, ca_s, h)[2], 2) if h > 19 else None])
for c in k[1]: c.font = Font(bold=True)
for col, w in zip('ABCDEFG', (32, 14, 11, 15, 15, 24, 24)): k.column_dimensions[col].width = w
raw = wb.create_sheet('Copernicus raw')
raw.append(['BaseLevel_m', 'Pixels', 'Area_m2', 'Capacity_m3', 'Area_ha', 'Area_km2', 'Capacity_MCM', 'Area_acre', 'Capacity_acre_ft'])
for r in eac_table(zc, ca, 21.0, 28.5, 0.25): raw.append(r)
rs = wb.create_sheet('SRTM raw')
rs.append(raw[1][0].value and [c.value for c in raw[1]])
for r in eac_table(zs, ca_s, 19.0, 28.5, 0.25): rs.append(r)
notes = wb.create_sheet('Notes')
for line in [
 ['Reservoir', 'Chembarambakkam Tank, Kancheepuram / Chennai (surplus regulator 12.9999 N, 80.0813 E)'],
 ['Prepared', '18 Oct 2026 with the area-elevation-capacity skill; planning level'],
 ['DEM', 'Copernicus GLO-30 (2011-2015, EGM2008), reprojected to UTM 44N at 30 m; SRTM 1" as cross-check'],
 ['Polygon', 'chembarambakkam_area_28.5m.geojson: cells below 28.5 m connected to the lake, bund axis burnt as a wall (traced on the DEM crest)'],
 ['Basin closure', 'Holds to 28.5 m. At ~28.6 m the DEM spills over a low point near 13.059 N, 80.053 E (north of the lake), below TBL 28.77 m. Check on the ground: real saddle, bund or DSM artefact'],
 ['Below 21.0 m', f'Copernicus flattens the lake at 21.0 m ({A21:.2f} km2), so storage below it is invisible. Filled with a power law from gauge zero {GZ:.2f} m: S = {S21:.2f} MCM x ((h-{GZ:.2f})/{FLAT-GZ:.2f})^{b:.2f}, with S(21.0) chosen so S(FRL) = published 3,645 Mcft ({S_PUB:.2f} MCM)'],
 ['Levels (WRD)', f'FRL {FRL:.2f} m (85.40 ft), MWL {MWL:.2f} m, TBL {TBL:.2f} m, gauge zero {GZ:.2f} m (FRL - 24 ft), sill {SILL:.2f} m (from the hydrology-dambreak study)'],
 ['Datum', 'Copernicus heights are EGM2008; WRD levels are GTS/MSL. Agreement is supported by the DEM bund crest (27.2-28.9 m vs TBL 28.77 m) but not proven; an offset of +/-1 m would shift every level'],
 ['Capacity', '3,645 Mcft is the design gross capacity. Sedimentation has reportedly reduced it; a bathymetric survey would replace the power-law part'],
 ['SRTM', 'SRTM (Feb 2000) flattens the lake at 19 m, has artefacts down to -1.9 m and gives more storage above 19 m than the published total; not used'],
]: notes.append(line)
notes.column_dimensions['A'].width = 16; notes.column_dimensions['B'].width = 150
wb.save(OUT + 'Chembarambakkam_EAC.xlsx')
# chart
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
H = np.arange(GZ, 28.501, 0.05); AD = np.array([adopted(h)[:2] for h in H])
fig, ax = plt.subplots(figsize=(10, 7))
ax.plot(AD[:, 1], H, color='#1f5aa6', lw=2.5, label='Capacity – adopted (MCM)')
ax.plot(st_s[st_h <= 28.5], st_h[st_h <= 28.5], '--', color='#1f5aa6', alpha=.6, lw=1.2, label='Capacity – study assumed curve')
Hs = H[H > 19]; ax.plot([dem(zs, ca_s, h)[2] for h in Hs], Hs, ':', color='grey', lw=1.2, label='Capacity – SRTM from 19 m (not used)')
ax.set_xlabel('Capacity (MCM)', color='#1f5aa6'); ax.set_ylabel('Water level (m)'); ax.set_xlim(0, 200); ax.set_ylim(18.5, 29.2); ax.grid(alpha=.3)
top = ax.twiny()
top.plot(AD[:, 0], H, color='#c8551b', lw=2.5, label='Area – adopted (km²)')
top.plot(st_a[st_h <= 28.5], st_h[st_h <= 28.5], '--', color='#c8551b', alpha=.6, lw=1.2, label='Area – study assumed curve')
top.set_xlim(40, 0); top.set_xlabel('Area (km²)', color='#c8551b')
for name, h, c in (('Gauge zero', GZ, 'k'), ('Sill', SILL, 'k'), ('FRL', FRL, 'g'), ('MWL', MWL, 'orange'), ('TBL', TBL, 'r'), ('DEM water surface', FLAT, 'grey')):
    ax.axhline(h, color=c, lw=0.8, ls='-.'); ax.text(198, h + 0.05, f'{name} {h:.2f} m', ha='right', fontsize=8, color=c)
ax.axhspan(28.5, 29.2, color='red', alpha=0.06); ax.text(5, 28.85, 'above 28.5 m: DEM basin spills north', fontsize=8, color='r')
h1, l1 = ax.get_legend_handles_labels(); h2, l2 = top.get_legend_handles_labels()
fig.legend(h1 + h2, l1 + l2, loc='lower center', ncol=3, fontsize=8, frameon=False)
fig.suptitle('Chembarambakkam Tank – Elevation-Area-Capacity Curves (Copernicus GLO-30, planning level)')
fig.tight_layout(rect=(0, 0.07, 1, 1)); fig.savefig(OUT + 'Chembarambakkam_EAC.png', dpi=150)
print(f'A21 {A21:.3f} km2  S21 {S21:.2f} MCM  b {b:.3f}')
for name, h in key:
    a, s, _ = adopted(h)
    print(f'{name:32s} {h:6.2f} m  {a:6.2f} km2  {s:7.2f} MCM  {s/MCFT:6.0f} Mcft  study {np.interp(h, st_h, st_s):7.2f}')
